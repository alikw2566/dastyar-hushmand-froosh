"""Real PDF and multi-sheet XLSX exports generated from persisted data."""

from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import Any

import arabic_reshaper
from bidi.algorithm import get_display
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _value(item: Any, key: str, default=None):
    return item.get(key, default) if isinstance(item, dict) else getattr(item, key, default)


def _find_font(configured: str = "") -> Path:
    candidates = [
        Path(configured) if configured else None,
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/tahoma.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate and candidate.is_file():
            return candidate
    raise RuntimeError("PDF_FONT_PATH must point to a Unicode Persian-capable TTF font")


def _rtl(value: object) -> str:
    text = "—" if value in (None, "") else str(value)
    return get_display(arabic_reshaper.reshape(text))


def render_call_pdf(
    call: Any,
    segments: Iterable[Any],
    analysis: dict | None,
    *,
    font_path: str = "",
) -> bytes:
    output = BytesIO()
    font = _find_font(font_path)
    font_name = f"Persian-{abs(hash(str(font))) & 0xFFFF:x}"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font)))
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Sales call report",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PersianTitle", parent=styles["Title"], fontName=font_name, alignment=TA_RIGHT, fontSize=16
    )
    body_style = ParagraphStyle(
        "PersianBody", parent=styles["BodyText"], fontName=font_name, alignment=TA_RIGHT, leading=18
    )
    story = [Paragraph(_rtl("گزارش تحلیل تماس فروش"), title_style), Spacer(1, 5 * mm)]
    summary_rows = [
        [_rtl("نام فایل"), _rtl(_value(call, "original_file_name"))],
        [_rtl("مشتری"), _rtl(_value(call, "customer_name"))],
        [_rtl("فروشنده"), _rtl(_value(call, "seller_name"))],
        [_rtl("نتیجه"), _rtl(_value(call, "outcome"))],
        [_rtl("امتیاز"), _rtl(_value(call, "score"))],
        [_rtl("مدت تماس (ثانیه)"), _rtl(_value(call, "duration_seconds"))],
    ]
    table = Table(summary_rows, colWidths=[48 * mm, 112 * mm], hAlign="RIGHT")
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, "#CBD5E1"),
                ("BACKGROUND", (0, 0), (0, -1), "#E2E8F0"),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([table, Spacer(1, 5 * mm)])
    if analysis:
        story.append(Paragraph(_rtl("خلاصه مدیریتی"), title_style))
        story.append(Paragraph(_rtl(analysis.get("executive_summary")), body_style))
        story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(_rtl("متن کامل مکالمه"), title_style))
    for segment in segments:
        role = _value(segment, "speaker_role", _value(segment, "role", "unknown"))
        content = _value(segment, "content", _value(segment, "text", ""))
        start = _value(segment, "start_seconds", _value(segment, "start"))
        prefix = f"[{float(start):.1f}] " if start is not None else ""
        story.append(Paragraph(_rtl(f"{prefix}{role}: {content}"), body_style))
    document.build(story)
    return output.getvalue()


def render_calls_xlsx(calls: Iterable[dict], *, title: str = "تماس‌ها") -> bytes:
    rows = list(calls)
    workbook = Workbook()
    calls_sheet = workbook.active
    calls_sheet.title = "Calls"
    headers = [
        "id",
        "created_at",
        "customer_name",
        "seller_name",
        "company",
        "phone",
        "city",
        "province",
        "product",
        "product_category",
        "outcome",
        "sales_stage",
        "lead_temperature",
        "score",
        "duration_seconds",
        "status",
        "followup_required",
        "followup_due_at",
        "error_code",
    ]
    calls_sheet.append(headers)
    for row in rows:
        calls_sheet.append([row.get(header) for header in headers])
    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Report", title])
    metadata.append(["Rows", len(rows)])
    metadata.append(["Generated by", "Mokalemeban"])
    for sheet in workbook.worksheets:
        sheet.sheet_view.rightToLeft = True
        sheet.freeze_panes = "A2" if sheet is calls_sheet else None
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                45, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_call_xlsx(
    call: dict, segments: list[dict], analysis: dict | None, tasks: list[dict]
) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Call"
    for key, value in call.items():
        overview.append([key, value])
    transcript = workbook.create_sheet("Transcript")
    transcript.append(["position", "start", "end", "speaker", "role", "content", "corrected"])
    for item in segments:
        transcript.append(
            [
                item.get(key)
                for key in (
                    "position",
                    "start",
                    "end",
                    "speaker",
                    "role",
                    "content",
                    "manually_corrected",
                )
            ]
        )
    analysis_sheet = workbook.create_sheet("Analysis")
    analysis_sheet.append(["field", "value"])
    for key, value in (analysis or {}).items():
        analysis_sheet.append([key, str(value)])
    followups = workbook.create_sheet("FollowUps")
    followups.append(["title", "priority", "status", "due_at"])
    for task in tasks:
        followups.append([task.get(key) for key in ("title", "priority", "status", "due_at")])
    for sheet in workbook.worksheets:
        sheet.sheet_view.rightToLeft = True
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
