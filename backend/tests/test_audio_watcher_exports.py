from __future__ import annotations

import wave
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.services.audio import preprocess_audio, validate_audio
from app.services.exports import render_call_pdf, render_call_xlsx, render_calls_xlsx
from app.services.pipeline import PipelineFailure
from app.services.watcher import (
    DurableObservation,
    ImportResult,
    LocalRecordingWatcher,
    source_file_name,
)


def make_wav(path: Path, seconds: float = 2.0) -> None:
    frames = int(16_000 * seconds)
    with wave.open(str(path), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(16_000)
        output.writeframes(b"\0\0" * frames)


def test_sftp_source_preserves_original_recording_name(tmp_path):
    staged = tmp_path / "random-prefix-call-123.wav"
    assert (
        source_file_name(staged, "sftp://pbx.example:22/recordings/2026/09/call-123.wav")
        == "call-123.wav"
    )
    assert source_file_name(staged, str(staged)) == staged.name


def test_audio_validation_and_standard_preprocessing(tmp_path):
    source = tmp_path / "call.wav"
    output = tmp_path / "processed.wav"
    make_wav(source)
    metadata = validate_audio(
        source,
        allowed_extensions=frozenset({".wav"}),
        max_bytes=10_000_000,
        min_duration_seconds=1,
    )
    assert metadata.duration_seconds == 2
    result = preprocess_audio(source, output, metadata)
    assert result.copied_without_conversion is True
    assert result.metadata_after.sample_rate == 16_000
    short = tmp_path / "short.wav"
    make_wav(short, 0.1)
    with pytest.raises(PipelineFailure, match="audio is too short"):
        validate_audio(
            short, allowed_extensions=frozenset({".wav"}), max_bytes=100_000, min_duration_seconds=1
        )


@pytest.mark.asyncio
async def test_watcher_import_duplicate_and_heartbeat(tmp_path):
    recording = tmp_path / "in-09121234567-200-20260801-120000-id.wav"
    make_wav(recording)

    class Repository:
        def __init__(self):
            self.imported = []
            self.heartbeats = []
            self.duplicate = False

        async def observe(self, source_identifier, path, observed_at, stability_seconds):
            return DurableObservation(True, datetime(2026, 8, 1, tzinfo=UTC))

        async def is_duplicate(self, source_identifier, sha256):
            return self.duplicate

        async def import_file(self, path, source_identifier, sha256, detected_at, metadata):
            self.imported.append((path, sha256, detected_at, metadata))
            return ImportResult("imported", "call-id")

        async def quarantine(self, *args):
            raise AssertionError("valid WAV must not be quarantined")

        async def heartbeat(self, status, metrics, error=None):
            self.heartbeats.append((status, metrics, error))

    repository = Repository()
    settings = Settings(
        _env_file=None,
        issabel_recordings_path=str(tmp_path),
        issabel_quarantine_path=str(tmp_path / "quarantine"),
        issabel_file_stability_seconds=0,
        issabel_allowed_extensions=".wav",
    )
    watcher = LocalRecordingWatcher(settings, repository)
    first = await watcher.scan_once()
    assert first["imported"] == 1
    assert repository.imported[0][3].direction == "inbound"
    assert repository.heartbeats[-1][0] == "healthy"
    repository.duplicate = True
    second = await watcher.scan_once()
    assert second["duplicates"] == 1


def test_watcher_rejects_symlink_that_escapes_drop_folder(tmp_path):
    root = tmp_path / "drop"
    root.mkdir()
    outside = tmp_path / "outside.wav"
    make_wav(outside)
    link = root / "linked.wav"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("creating symlinks is not permitted on this Windows host")

    class Repository:
        pass

    settings = Settings(
        _env_file=None,
        issabel_recordings_path=str(root),
        issabel_quarantine_path=str(root / "quarantine"),
        issabel_allowed_extensions=".wav",
    )
    watcher = LocalRecordingWatcher(settings, Repository())
    assert watcher._candidate_paths() == []


def test_pdf_and_multisheet_excel_are_real_files():
    call = {
        "id": "call-1",
        "original_file_name": "sample.wav",
        "customer_name": "مشتری آزمون",
        "seller_name": "کارشناس",
        "outcome": "follow_up",
        "score": 82,
        "duration_seconds": 65,
    }
    segments = [
        {
            "position": 0,
            "start": 0,
            "end": 2,
            "speaker": "A",
            "role": "agent",
            "content": "سلام",
            "manually_corrected": False,
        }
    ]
    pdf = render_call_pdf(call, segments, {"executive_summary": "خلاصه فارسی"})
    assert pdf.startswith(b"%PDF-")
    assert len(pdf) > 1_000
    single = render_call_xlsx(
        call, segments, {"executive_summary": "summary"}, [{"title": "پیگیری", "status": "open"}]
    )
    workbook = load_workbook(BytesIO(single))
    assert workbook.sheetnames == ["Call", "Transcript", "Analysis", "FollowUps"]
    multiple = render_calls_xlsx([call])
    workbook = load_workbook(BytesIO(multiple))
    assert workbook.sheetnames == ["Calls", "Metadata"]
    assert workbook["Calls"].max_row == 2
